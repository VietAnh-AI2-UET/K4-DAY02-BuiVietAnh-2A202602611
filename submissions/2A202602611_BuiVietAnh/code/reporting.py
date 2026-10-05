"""Bước 5: kiểm tra số liệu, hoàn thiện Excel, tạo báo cáo và bộ file nộp.

Chỉ đọc log/CSV đã có. Không huấn luyện, không gọi model trên test lần nữa.
"""
from __future__ import annotations

from pathlib import Path
import hashlib
import json
import shutil
import zipfile
import numpy as np
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import dataset
import experiments
import train
from eval import check_against_csv, compute_metrics, load_group, read_pred, save_group

SHEETS = ("Backbones", "Training", "Inference", "Final", "PerClass", "Latency", "Summary")
TOLERANCE = 2e-6  # CSV làm tròn xác suất; cho phép sai số rất nhỏ khi tính ECE.


def require_close(actual, expected, description):
    if not np.isfinite(actual) or not np.isclose(float(actual), float(expected), atol=TOLERANCE, rtol=0):
        raise ValueError(f"Số liệu không khớp: {description}: bảng={actual}, tính lại={expected}")


def metrics_from_file(path, reference, split):
    prediction = read_pred(str(path))
    check_against_csv(prediction, str(reference), split)
    metrics = compute_metrics(prediction.y_true, prediction.y_pred, prediction.probs)
    return prediction, metrics


def verify_tables(tables, runs, predictions, labels, fold, seeds):
    """Đối chiếu từng dòng với CSV và checkpoint/log của đúng exp_id, seed."""
    val_csv, test_csv = labels / f"val_subset{fold}.csv", labels / f"test_subset{fold}.csv"
    evidence = {}
    for sheet in ("Backbones", "Training", "Final"):
        frame = tables[sheet]
        if sheet == "Final":
            frame = frame.loc[frame.kind.eq("seed")]
        for row in frame.to_dict("records"):
            tag, seed = row["exp_id"], int(row["seed"])
            folder = runs / tag / f"seed{seed}"
            for name in ("config.json", "history.csv", "summary.json", "best.pt"):
                if not (folder / name).is_file():
                    raise FileNotFoundError(f"Thiếu bằng chứng lần chạy: {folder / name}")
            config, summary = experiments.read_json(folder / "config.json"), experiments.read_json(folder / "summary.json")
            if config["exp_id"] != tag or config["seed"] != seed:
                raise ValueError(f"Cấu hình không thuộc lần chạy {tag}, seed {seed}")
            if row.get("backbone") != config["backbone"] or row.get("weight_tag") != summary["weight_tag"]:
                raise ValueError(f"{tag} seed{seed}: backbone/tag trong bảng không khớp lần chạy")
            history = pd.read_csv(folder / "history.csv")
            if len(history) != int(config["epochs"]) or history.epoch.tolist() != list(range(1, int(config["epochs"]) + 1)):
                raise ValueError(f"{folder}: history thiếu epoch hoặc sai thứ tự")
            # idxmax lấy epoch sớm nhất khi điểm hoà, đúng quy tắc của train.run.
            best = history.loc[history.val_macro_f1.idxmax()]
            require_close(summary["val_macro_f1"], best.val_macro_f1, f"{tag} checkpoint tốt nhất")
            if int(summary["best_epoch"]) != int(best.epoch):
                raise ValueError(f"{tag}: best_epoch không khớp history")
            require_close(summary["train_seconds_per_epoch"], history.train_seconds.mean(), f"{tag} thời gian/epoch")
            for key in ("params_M", "GMAC", "train_seconds_per_epoch"):
                if key in row:
                    require_close(row[key], summary[key], f"{tag} seed{seed} {key}")
            _, metrics = metrics_from_file(predictions / f"{tag}_seed{seed}_val.csv", val_csv, "val")
            for column, key in (("val_macro_f1", "macro_f1"), ("val_top1", "top1")):
                require_close(row[column], metrics[key], f"{sheet}/{tag} seed{seed} {column}")
            evidence[(tag, seed)] = (config, history)

    inference = tables["Inference"]
    inference_seed = int(experiments.read_json("training_selection.json")["best_config"]["seed"])
    for row in inference.to_dict("records"):
        _, metrics = metrics_from_file(predictions / f"{row['exp_id']}_seed{inference_seed}_val.csv", val_csv, "val")
        for column, key in (("val_macro_f1", "macro_f1"), ("val_top1", "top1"), ("val_ece", "ece"), ("val_nll", "nll")):
            require_close(row[column], metrics[key], f"{row['exp_id']} {column}")

    groups = {}
    final_seed_rows = tables["Final"].loc[tables["Final"].kind.eq("seed")]
    for tag in ("F01", "T00"):
        group = load_group([str(predictions / f"{tag}_seed{seed}_test.csv") for seed in seeds], str(test_csv))
        groups[tag] = group
        subset = final_seed_rows.loc[final_seed_rows.exp_id.eq(tag)]
        if sorted(subset.seed.astype(int).tolist()) != sorted(seeds):
            raise ValueError(f"{tag}: số seed trong bảng khác cấu hình chung kết đã chốt")
        for prediction, metrics in zip(group.preds, group.metrics):
            row = subset.loc[subset.seed.astype(int).eq(prediction.seed)].iloc[0]
            for column, key in (("test_macro_f1", "macro_f1"), ("test_top1", "top1"), ("test_ece", "ece")):
                require_close(row[column], metrics[key], f"{tag} seed{prediction.seed} {column}")
            class_rows = tables["PerClass"]
            class_rows = class_rows.loc[class_rows.exp_id.eq(tag) & class_rows.seed.astype(str).eq(str(prediction.seed))]
            if len(class_rows) != dataset.NUM_CLASSES:
                raise ValueError(f"{tag} seed{prediction.seed}: bảng PerClass thiếu lớp")
            for label, name in enumerate(dataset.CLASS_NAMES):
                row_class = class_rows.loc[class_rows["class"].eq(name)].iloc[0]
                for key in ("precision", "recall", "f1", "support"):
                    require_close(row_class[key], metrics[key][label], f"{tag}/{name} {key}")
        aggregates = tables["Final"].loc[tables["Final"].exp_id.eq(tag) & tables["Final"].kind.eq("mean_std")]
        if len(aggregates) != 1:
            raise ValueError(f"{tag}: thiếu dòng mean/std trong bảng Final")
        for column, key in (("test_macro_f1", "macro_f1"), ("test_top1", "top1"), ("test_ece", "ece")):
            require_close(aggregates.iloc[0][column], group.summary[key][0], f"{tag} mean {key}")
            require_close(aggregates.iloc[0][column + "_std"], group.summary[key][1], f"{tag} std {key}")
        for column in ("val_macro_f1", "val_top1", "p95_ms"):
            require_close(aggregates.iloc[0][column], subset[column].mean(), f"{tag} mean {column}")
            require_close(aggregates.iloc[0][column + "_std"], subset[column].std(ddof=1), f"{tag} std {column}")
        for column in ("test_ece_uncal",):
            require_close(aggregates.iloc[0][column], subset[column].mean(), f"{tag} mean {column}")
            require_close(aggregates.iloc[0][column + "_std"], subset[column].std(ddof=1), f"{tag} std {column}")
        for row in subset.to_dict("records"):
            require_close(row["val_test_f1_gap"], abs(row["val_macro_f1"] - row["test_macro_f1"]), f"{tag} val/test gap")
        class_averages = tables["PerClass"].loc[tables["PerClass"].exp_id.eq(tag)
                                              & tables["PerClass"].seed.astype(str).eq("mean_std")]
        if len(class_averages) != 9:
            raise ValueError(f"{tag}: thiếu trung bình từng lớp")
        for label, name in enumerate(dataset.CLASS_NAMES):
            average = class_averages.loc[class_averages["class"].eq(name)].iloc[0]
            require_close(average.support, group.metrics[0]["support"][label], f"{tag}/{name} mean support")
            for key in ("precision", "recall", "f1"):
                require_close(average[key], group.summary[key][0][label], f"{tag}/{name} mean {key}")
                require_close(average[key + "_std"], group.summary[key][1][label], f"{tag}/{name} std {key}")
    uncal = load_group([str(predictions / f"F01uncal_seed{seed}_test.csv") for seed in seeds], str(test_csv))
    groups["F01uncal"] = uncal
    for prediction, metrics in zip(uncal.preds, uncal.metrics):
        row = final_seed_rows.loc[final_seed_rows.exp_id.eq("F01") & final_seed_rows.seed.astype(int).eq(prediction.seed)].iloc[0]
        require_close(row.test_ece_uncal, metrics["ece"], f"F01 seed{prediction.seed} ECE trước calibration")
    latency = tables["Latency"]
    # Loại trừ độ trễ sơ bộ bước 1; kiểm tra điều kiện đo chi tiết bước 3–4.
    if latency.empty or not {1, 32}.issubset(set(latency.batch.astype(int))):
        raise ValueError("Bảng Latency phải có batch 1 và batch 32")
    if (latency.n < 50).any() or (latency.p50 <= 0).any() or (latency.p95 < latency.p50).any() or (latency.p99 < latency.p95).any():
        raise ValueError("Độ trễ chưa đủ 50 lần đo hoặc số percentile không hợp lệ")
    for row in inference.to_dict("records"):
        measured = latency.loc[latency.exp_id.eq(row["exp_id"]) & latency.batch.eq(1)]
        if len(measured) != 1:
            raise ValueError(f"{row['exp_id']}: thiếu hoặc trùng phép đo batch 1")
        for key in ("p50", "p95", "p99"):
            require_close(row[key + "_ms"], measured.iloc[0][key], f"{row['exp_id']} {key}")
    for row in final_seed_rows.to_dict("records"):
        measured = latency.loc[latency.exp_id.eq(row["exp_id"]) & latency.batch.eq(1)
                               & pd.to_numeric(latency.seed, errors="coerce").eq(int(row["seed"]))]
        if len(measured) != 1:
            raise ValueError(f"{row['exp_id']} seed{row['seed']}: thiếu hoặc trùng độ trễ chung kết")
        require_close(row["p95_ms"], measured.iloc[0].p95, f"{row['exp_id']} seed{row['seed']} p95")
    return evidence, groups


def build_summary(tables):
    """Top 10 theo VAL; mỗi cấu hình một dòng, vòng cuối lấy trung bình qua seed."""
    rows = []
    for sheet in ("Backbones", "Training"):
        for row in tables[sheet].to_dict("records"):
            if row["exp_id"] == "T00":
                continue  # T00 có đủ seed ở bảng Final; không lặp dòng một seed ở đây.
            rows.append({**{key: row.get(key, np.nan) for key in
                            ("exp_id", "backbone", "weight_tag", "val_macro_f1", "val_top1", "params_M", "GMAC", "train_seconds_per_epoch")},
                         "source": sheet, "n_seeds": 1, "p95_ms": np.nan,
                         "latency_batch1_ms": row.get("latency_batch1_ms", np.nan)})
    selected = experiments.read_json("training_selection.json")
    trained = tables["Training"].loc[tables["Training"].exp_id.eq(selected["best_exp_id"])].iloc[0]
    for row in tables["Inference"].to_dict("records"):
        rows.append({**row, "source": "Inference", "n_seeds": 1, "backbone": trained.backbone,
                     "params_M": trained.params_M, "GMAC": trained.GMAC * row["K"],
                     "weight_tag": trained.weight_tag, "train_seconds_per_epoch": np.nan})
    final = tables["Final"]
    for tag in ("F01", "T00"):
        mean = final.loc[final.exp_id.eq(tag) & final.kind.eq("mean_std")].iloc[0]
        seeds = final.loc[final.exp_id.eq(tag) & final.kind.eq("seed")]
        rows.append({"exp_id": tag, "source": "Final", "backbone": seeds.iloc[0].backbone,
                     "weight_tag": seeds.iloc[0].weight_tag, "n_seeds": len(seeds),
                     "val_macro_f1": mean.val_macro_f1, "val_macro_f1_std": mean.val_macro_f1_std,
                     "val_top1": mean.val_top1, "params_M": seeds.params_M.mean(),
                     "GMAC": seeds.GMAC.mean() * {"identity": 1, "hflip_prob": 2, "hflip_logit": 2, "fivecrop": 5}[seeds.iloc[0].method],
                     "p95_ms": seeds.p95_ms.mean(), "train_seconds_per_epoch": seeds.train_seconds_per_epoch.mean()})
    frame = pd.DataFrame(rows).sort_values(["val_macro_f1", "exp_id"], ascending=[False, True]).head(10).reset_index(drop=True)
    frame.insert(0, "rank", range(1, len(frame) + 1))
    return frame[["rank", "exp_id", "source", "backbone", "n_seeds", "val_macro_f1", "val_macro_f1_std",
                  "val_top1", "params_M", "GMAC", "train_seconds_per_epoch", "latency_batch1_ms", "p95_ms"]].rename(columns={"GMAC": "GMAC_estimate"})


def format_workbook(workbook, tables):
    temporary = workbook.with_name(workbook.stem + ".step5.tmp.xlsx")
    shutil.copy2(workbook, temporary)
    # Giữ các sheet riêng của người dùng; chỉ cập nhật 7 sheet của bài lab.
    with pd.ExcelWriter(temporary, engine="openpyxl", mode="a", if_sheet_exists="replace") as writer:
        for name in SHEETS:
            frame = tables[name]
            frame.to_excel(writer, sheet_name=name, index=False)
            sheet = writer.sheets[name]
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            for cell in sheet[1]:
                cell.fill = PatternFill("solid", fgColor="203864")
                cell.font = Font(color="FFFFFF", bold=True)
                cell.alignment = Alignment(wrap_text=True, vertical="center")
            sheet.row_dimensions[1].height = 32
            for index, column in enumerate(frame.columns, start=1):
                lengths = [len(str(column)), *[len(str(value)) for value in frame[column].head(100).fillna("")]]
                sheet.column_dimensions[get_column_letter(index)].width = min(48, max(12, max(lengths) + 2))
                for cells in sheet.iter_rows(min_row=2, min_col=index, max_col=index):
                    cell = cells[0]
                    if isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
                        cell.number_format = "0" if column in {"seed", "epoch", "epochs", "best_epoch", "rank", "n_seeds", "batch", "n", "K", "img_size", "support"} else "0.0000"
            if "val_macro_f1" in frame and frame.val_macro_f1.notna().any():
                best = int(frame.val_macro_f1.idxmax()) + 2
                for cell in sheet[best]:
                    cell.fill = PatternFill("solid", fgColor="E2F0D9")
    temporary.replace(workbook)


def markdown_table(frame):
    """Tự viết bảng Markdown để không cần cài thêm thư viện tabulate."""
    def cell(value):
        if pd.isna(value):
            return "—"
        if isinstance(value, (float, np.floating)):
            return f"{value:.4f}"
        return str(value).replace("|", "/").replace("\n", " ")
    lines = ["| " + " | ".join(map(str, frame.columns)) + " |", "| " + " | ".join(["---"] * len(frame.columns)) + " |"]
    lines += ["| " + " | ".join(cell(value) for value in row) + " |" for row in frame.itertuples(index=False, name=None)]
    return "\n".join(lines)


def error_gallery(group, images_dir, curves):
    """Xem ảnh đoán sai từ CSV đã lưu, không chạy lại model hoặc sửa cấu hình."""
    prediction = group.preds[0]
    wrong = np.flatnonzero(prediction.y_true != prediction.y_pred)
    if len(wrong) == 0:
        return "Seed đầu tiên không có ảnh đoán sai.", pd.DataFrame()
    # Ưu tiên ảnh model tự tin nhưng đoán sai, gồm cả hai lớp khó khi có.
    confidence = prediction.probs.max(1)
    ordered = sorted(wrong, key=lambda i: (prediction.y_true[i] not in {0, 7}, -confidence[i]))[:9]
    rows = [{"Filename": str(prediction.filenames[i]), "true": dataset.CLASS_NAMES[prediction.y_true[i]],
             "predicted": dataset.CLASS_NAMES[prediction.y_pred[i]], "confidence": float(confidence[i])} for i in ordered]
    table = pd.DataFrame(rows)
    from PIL import Image
    import matplotlib.pyplot as plt
    fig, axes = plt.subplots(3, 3, figsize=(12, 12))
    shown = 0
    for ax, row in zip(axes.flat, rows):
        image_path = images_dir / row["Filename"]
        if image_path.is_file():
            with Image.open(image_path) as image:
                ax.imshow(image.convert("RGB"))
            shown += 1
        else:
            ax.text(0.5, 0.5, "Không tìm thấy ảnh", ha="center")
        ax.set_title(f"Thật: {row['true']}\nĐoán: {row['predicted']} ({row['confidence']:.1%})", fontsize=9)
    for ax in axes.flat:
        ax.axis("off")
    fig.tight_layout()
    fig.savefig(curves / "F01_error_examples.png", dpi=150)
    plt.close(fig)
    return f"Minh hoạ {shown}/{len(rows)} ảnh đoán sai của seed {prediction.seed}; không dùng các ảnh này để chọn lại cấu hình.", table


def create_report(tables, groups, config, lock, errors_note, errors_table):
    final, baseline = groups["F01"], groups["T00"]
    delta = final.summary["macro_f1"][0] - baseline.summary["macro_f1"][0]
    noise = max(final.summary["macro_f1"][1], baseline.summary["macro_f1"][1])
    def metric(group, key):
        mean, std = group.summary[key]
        return f"{mean:.4f} ± {std:.4f}"
    label_dir = Path(config["labels_dir"])
    counts = {}
    for split in ("train", "val", "test"):
        df = pd.read_csv(label_dir / f"{split}_subset{config['fold']}.csv")
        counts[split] = df.Label.value_counts().reindex(range(9), fill_value=0).to_numpy()
    distribution = pd.DataFrame({"class": dataset.CLASS_NAMES, **counts})
    backbones, training, inference = tables["Backbones"], tables["Training"], tables["Inference"]
    best_backbone = backbones.loc[backbones.val_macro_f1.idxmax()]
    best_training = training.loc[training.val_macro_f1.idxmax()]
    combo = training.loc[training.exp_id.eq("T09")].iloc[0]
    latency = tables["Latency"]
    measured = latency.loc[latency.exp_id.eq("F01") & latency.batch.eq(1)]
    p95 = measured.p95.max()
    mean_rows = tables["Final"].loc[tables["Final"].kind.eq("mean_std")]
    before, after = groups["F01uncal"].summary["ece"][0], final.summary["ece"][0]
    confusion = sum(m["confusion"] for m in final.metrics)
    np.fill_diagonal(confusion, 0)
    pairs = sorted([(int(confusion[i, j]), dataset.CLASS_NAMES[i], dataset.CLASS_NAMES[j])
                    for i in range(9) for j in range(9) if i != j], reverse=True)[:5]
    pair_table = pd.DataFrame(pairs, columns=["count_across_seeds", "true", "predicted"])
    note = "Δ lớn hơn std quan sát; đây không phải kiểm định thống kê." if delta > noise else "Chưa có cải thiện vượt std quan sát."
    runtime = "Đạt ngưỡng p95 ≤ 100 ms trong điều kiện đã đo." if p95 <= 100 else "Chưa đạt ngưỡng p95 ≤ 100 ms; cần cải thiện tốc độ trước khi triển khai."
    per_class = tables["PerClass"].loc[tables["PerClass"].seed.astype(str).eq("mean_std")]
    recipe = {key: config[key] for key in ("backbone", "init", "aug", "loss", "sampler", "mix", "ema_decay", "img_size", "epochs", "batch_size", "lr_backbone", "lr_head", "weight_decay", "warmup_epochs")}
    lines = ["# Báo cáo Lab Day 2 — DeepWeeds", "", "## 1. Tóm tắt", "",
             f"So sánh {len(backbones)} backbone, {len(training)} cấu hình train và {len(inference)} cách suy luận trên val.",
             f"F01 trên test: macro-F1 {metric(final, 'macro_f1')}; top-1 {metric(final, 'top1')}.",
             f"Mốc T00: macro-F1 {metric(baseline, 'macro_f1')}; Δ={delta:+.4f}; std lớn hơn={noise:.4f}. {note}",
             f"ECE test trước/sau hiệu chỉnh: {before:.4f}/{after:.4f}. Số seed: {len(lock['seeds'])}.",
             "", "## 2. Dữ liệu và thiết lập", "", "Fold 0, giữ nguyên CSV chia sẵn; chỉ val được dùng để chọn cấu hình.",
             markdown_table(distribution), "", "Macro-F1 coi các lớp quan trọng như nhau; top-1 là tỉ lệ dự đoán đúng.",
             "ECE đo độ lệch giữa sự tự tin của dự đoán và tỉ lệ đúng; giá trị thấp hơn tốt hơn.",
             "", "Cấu hình chung kết:", "", "```json", json.dumps(recipe, ensure_ascii=False, indent=2), "```", "",
             "Phần cứng/điều kiện đo:", "", markdown_table(latency[["gpu", "dtype", "batch", "img_size", "torch", "preprocessing"]].drop_duplicates()),
             "", "## 3. So sánh backbone", "", markdown_table(backbones[["exp_id", "backbone", "weight_tag", "params_M", "GMAC", "val_macro_f1", "val_top1", "train_seconds_per_epoch", "latency_batch1_ms"]]),
             "", f"{best_backbone.exp_id} đạt macro-F1 val cao nhất {best_backbone.val_macro_f1:.4f}.",
             "GMAC tính bằng fvcore, không đếm mọi phép toán; độ trễ ở bước 1 mới là số đo sơ bộ.",
             "", "## 4. Công thức huấn luyện", "", "T00 là công thức nền trên backbone đã chọn; B01 là mốc ResNet của bước 1.",
             markdown_table(training[["exp_id", "axis", "changed", "val_macro_f1", "delta_vs_T00"]]), "",
             f"Cấu hình tốt nhất trên val là {best_training.exp_id}. Kết hợp T09 có Δ={combo.delta_vs_T00:+.4f} so với T00.",
             "Các thí nghiệm T01–T09 chỉ có một seed, nên chưa xác định được độ dao động riêng của từng thay đổi.",
             "Không dùng std của chung kết để khẳng định một kỹ thuật riêng lẻ thắng chắc chắn.",
             "", "## 5. Suy luận và tốc độ", "", markdown_table(inference[["exp_id", "description", "K", "val_macro_f1", "val_ece", "p95_ms", "relative_cost"]]),
             "", "![Chất lượng và độ trễ](curves/inference_tradeoff.png)", "",
             f"Kiểu view đã chốt: `{lock['method']}`. T được khớp riêng trên val của mỗi seed, rồi giữ cố định trên test.",
             "ECE val sau khớp T được đo trên chính tập khớp; hiệu quả trên dữ liệu mới được kiểm tra bằng ECE test.",
             "", "## 6. Chung kết và phân tích lỗi", "", markdown_table(mean_rows[["exp_id", "val_macro_f1", "val_macro_f1_std", "test_macro_f1", "test_macro_f1_std", "test_top1", "test_top1_std", "test_ece"]]),
             "", markdown_table(per_class[["exp_id", "class", "support", "precision", "recall", "f1", "f1_std"]]),
             "", "![Ma trận nhầm lẫn F01](curves/F01_test_confusion.png)", "", markdown_table(pair_table),
             "", "Số nhầm lẫn ở bảng trên cộng qua các seed, không phải số ảnh khác nhau.", errors_note, "",
             "![Ảnh đoán sai](curves/F01_error_examples.png)" if not errors_table.empty else "Không có ảnh lỗi để minh hoạ.", "",
             markdown_table(errors_table) if not errors_table.empty else "",
             "Các cặp nhầm phổ biến gợi ý cần xem kỹ vùng cây và nền ảnh; chưa đủ bằng chứng để kết luận nguyên nhân sinh học.",
             "", "## 7. Kết luận", "", f"Δ F01 so với T00 = {delta:+.4f}. {note}",
             f"F01 có p95 lớn nhất qua seed = {p95:.2f} ms ở batch 1. {runtime}",
             f"ECE test {'giảm' if after < before else 'không giảm'} sau hiệu chỉnh ({before:.4f} → {after:.4f}).",
             "Chênh lệch qua backbone, recipe và cách suy luận chỉ là quan sát ở các điều kiện khác nhau; không cộng chúng như các đóng góp độc lập.",
             "", "## 8. Hạn chế", "", "Một fold; các vòng chọn cấu hình chỉ có một seed. Trọng số pretrained có thể có công thức huấn luyện khác nhau.",
             "Độ trễ không tính đọc/resize/normalize ảnh. Điều kiện phần cứng và tải máy có thể ảnh hưởng kết quả.",
             "Không thay đổi cấu hình sau khi xem test. Việc tiếp theo cần dùng dữ liệu mới hoặc một quy trình đánh giá độc lập.",
             "", "## 9. Phụ lục", "", "Top 10 cấu hình xếp theo val, không xếp theo test:", "", markdown_table(tables["Summary"]),
             "", "GMAC_estimate trong Summary là GMAC/forward nhân số view, chưa tính chi phí gộp; số đo sơ bộ latency_batch1_ms tách riêng p95_ms đo đủ 50 lần.",
             "", "Cấu hình/log/summary được chép vào `evidence/runs/`; kết quả tính lại bằng eval.py ở `eval_out/`.",
             "Các bản sao CSV, biểu đồ và code có mã SHA-256 trong `artifact_manifest.json` để kiểm tra file thay đổi.",
             "README hướng dẫn chạy lại; notebook ở `code/lab_day2.ipynb`.", ""]
    return "\n".join(lines)


def deliverables_step(workbook="results.xlsx", submission_dir=None):
    """Chạy sau bước 4. Chỉ xuất sản phẩm khi các kiểm tra số liệu đều đạt."""
    workbook = Path(workbook).resolve()
    tables = pd.read_excel(workbook, sheet_name=None)
    missing = [name for name in SHEETS[:-1] if name not in tables or tables[name].empty]
    if missing:
        raise ValueError(f"Chưa đủ kết quả để làm bước 5: thiếu sheet {missing}; hoàn tất bước 1–4 trước")
    if len(tables["Backbones"]) < 5 or len(tables["Inference"]) < 5 or "T09" not in set(tables["Training"].exp_id):
        raise ValueError("Thiếu backbone/cách suy luận/thí nghiệm kết hợp; hoàn tất các bước trước")
    selection = experiments.read_json("training_selection.json")
    config = selection["best_config"]
    runs, predictions, labels = map(Path, (config["out_dir"], config["pred_dir"], config["labels_dir"]))
    lock = experiments.read_json(runs / "final_selection_lock.json")
    inference_selection = experiments.read_json("inference_selection.json")
    if lock["final_config"] != config or lock["method"] != inference_selection["method"]:
        raise ValueError("Cấu hình chọn ở bước 2/3 không khớp cấu hình chung kết đã chốt")
    seeds = lock["seeds"]
    if len(seeds) < 3 or len(set(seeds)) != len(seeds) or config["fold"] != 0:
        raise ValueError("Cần fold 0 và ít nhất 3 seed khác nhau")
    evidence, groups = verify_tables(tables, runs, predictions, labels, config["fold"], seeds)
    tables["Summary"] = build_summary(tables)
    # Thêm mô tả đầy đủ và mean ± std dễ đọc, vẫn giữ các cột số để kiểm tra lại.
    final = tables["Final"]
    final["configuration"] = [json.dumps({key: evidence[(row.exp_id, int(row.seed))][0][key]
                                          for key in ("backbone", "aug", "loss", "sampler", "mix", "ema_decay")},
                                         ensure_ascii=False) + f"; inference={row.method}"
                              if row.kind == "seed" else "Trung bình ± std qua seed" for row in final.itertuples()]
    final["test_macro_f1_mean_std"] = [f"{row.test_macro_f1:.4f} ± {row.test_macro_f1_std:.4f}"
                                       if row.kind == "mean_std" else "" for row in final.itertuples()]
    curves = Path("curves")
    curves.mkdir(exist_ok=True)
    # Vẽ lại từ history, mỗi seed một file để không bị ghi đè giữa các seed.
    for (tag, seed), (run_config, history) in evidence.items():
        train.plot_curves(history.to_dict("records"), curves / f"{tag}_seed{seed}_{run_config['backbone']}.png",
                          f"{tag}, seed {seed} — {run_config['backbone']}")
    matrices = [(tag, p.seed, m["confusion"]) for tag in ("F01", "T00")
                for p, m in zip(groups[tag].preds, groups[tag].metrics)]
    experiments.plot_final_matrices(matrices)
    if not (curves / "inference_tradeoff.png").is_file():
        raise FileNotFoundError("Thiếu curves/inference_tradeoff.png; chạy bước 3 trước")
    errors_note, errors_table = error_gallery(groups["F01"], Path(config["images_dir"]), curves)
    format_workbook(workbook, tables)
    destination = Path(submission_dir).resolve() if submission_dir else Path(__file__).resolve().parent.parent
    destination.mkdir(parents=True, exist_ok=True)
    def copy_file(source, target):
        source, target = Path(source).resolve(), Path(target).resolve()
        target.parent.mkdir(parents=True, exist_ok=True)
        if source != target:
            shutil.copy2(source, target)
    copy_file(workbook, destination / "results.xlsx")
    for folder, target_name in ((curves, "curves"), (predictions, "predictions")):
        for source in folder.glob("*"):
            if source.is_file() and source.suffix.lower() in {".png", ".csv"}:
                copy_file(source, destination / target_name / source.name)
    code_dir = Path(__file__).resolve().parent
    for source in code_dir.iterdir():
        if source.suffix in {".py", ".ipynb"}:
            copy_file(source, destination / "code" / source.name)
    for tag, seed in evidence:
        for name in ("config.json", "history.csv", "summary.json"):
            copy_file(runs / tag / f"seed{seed}" / name, destination / "evidence/runs" / tag / f"seed{seed}" / name)
    evaluation = destination / "eval_out"
    for tag, group in groups.items():
        save_group(evaluation, tag, group, dataset.CLASS_NAMES)
    # Lưu kết quả đối chiếu dưới dạng máy đọc được; không khẳng định kết quả trước khi chạy.
    experiments.write_json(destination / "verification.json", {
        "status": "passed", "tolerance": TOLERANCE, "seeds": seeds,
        "sheets": list(SHEETS), "verified_runs": len(evidence),
        "note": "Tính lại từ CSV đã lưu bằng eval.py, không gọi model trên test."})
    report = create_report(tables, groups, config, lock, errors_note, errors_table)
    # Giữ bản người dùng đã chỉnh; luôn cập nhật bản generated để xem số mới nhất.
    (destination / "report.generated.md").write_text(report, encoding="utf-8")
    if not (destination / "report.md").exists():
        (destination / "report.md").write_text(report, encoding="utf-8")
    readme = "\n".join([
        "# Hướng dẫn chạy lại bài nộp", "", "Notebook: [code/lab_day2.ipynb](code/lab_day2.ipynb).", "",
        "Giữ nguyên vị trí bài nộp trong repo gốc chứa eval.py. Không sửa eval.py.",
        "Mở notebook bằng môi trường có PyTorch/torchvision phù hợp GPU; chạy ô cài thư viện và tải dữ liệu.",
        "Đặt IMAGES_DIR/LABELS_DIR đúng vị trí, chạy bước 0 → 1 → 2 → 3 → 4 → 5.",
        "Khi tiếp tục cùng kernel, không cần train lại bước 0/1. T00 là nền trên backbone được chọn ở bước 1.",
        "Đóng results.xlsx trong Excel trước khi chạy. Checkpoint và dữ liệu lớn không nằm trong bộ file nộp.",
        "submission_bundle.zip chỉ gồm file sản phẩm được liệt kê trong manifest, không chứa code/data hoặc code/runs.",
        "Cấu hình/version thư viện của từng lần chạy: evidence/runs/<exp_id>/seed<k>/config.json.",
        "Bước 4 giữ cấu hình cố định, mỗi cấu hình/seed đi qua test một lượt; chạy lại đọc cache.",
        "Bước 5 chỉ đọc CSV/log để kiểm tra và xuất sản phẩm. Đọc verification.json và eval_out/ để đối chiếu.",
        "report.generated.md là bản tự tạo mới nhất; report.md có thể chỉnh thủ công và sẽ không bị ghi đè.",
        "Trước khi nộp, xem lại ảnh lỗi/nhận xét và thêm link Colab/Kaggle nếu lớp yêu cầu chạy trên đó.", ""])
    (destination / "README.generated.md").write_text(readme, encoding="utf-8")
    if not (destination / "README.md").exists():
        (destination / "README.md").write_text(readme, encoding="utf-8")
    # Manifest chỉ ghi file nộp, không quét ảnh dữ liệu hoặc checkpoint hàng trăm MB.
    files = [destination / name for name in ("results.xlsx", "report.md", "report.generated.md", "README.md", "README.generated.md", "verification.json")]
    for name in ("curves", "predictions", "evidence", "eval_out"):
        files += [path for path in (destination / name).rglob("*") if path.is_file()]
    files += [path for path in (destination / "code").iterdir() if path.suffix in {".py", ".ipynb"}]
    manifest = {str(path.relative_to(destination)).replace("\\", "/"): hashlib.sha256(path.read_bytes()).hexdigest() for path in files}
    experiments.write_json(destination / "artifact_manifest.json", manifest)
    # Nén đúng danh sách sản phẩm, không nén toàn bộ thư mục code chứa data/runs.
    archive = destination / "submission_bundle.zip"
    with zipfile.ZipFile(archive, "w", compression=zipfile.ZIP_DEFLATED) as bundle:
        for path in [*files, destination / "artifact_manifest.json"]:
            bundle.write(path, arcname=str(path.relative_to(destination)))
    print("Bước 5 hoàn tất. Bộ file nộp:", destination)
    print("Đủ 7 sheet; số liệu khớp CSV/log. Không huấn luyện hoặc dự đoán test thêm.")
    print("Xem report.generated.md; report.md/README.md đã tồn tại được giữ nguyên.")
    print("File nén để nộp (không có dữ liệu/checkpoint lớn):", archive)
    return {"submission_dir": str(destination), "archive": str(archive), "summary": tables["Summary"], "verified_runs": len(evidence)}
