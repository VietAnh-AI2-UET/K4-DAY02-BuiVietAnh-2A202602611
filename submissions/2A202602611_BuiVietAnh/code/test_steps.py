"""Kiểm tra CPU bằng model/dữ liệu nhỏ; không đụng dữ liệu hoặc GPU đang train.

Chạy: python -m unittest test_steps -v (từ thư mục code).
"""
from pathlib import Path
from dataclasses import asdict, replace
from unittest.mock import patch
import os
import tempfile
import unittest
import numpy as np
import pandas as pd
import torch
from torch import nn

import losses
import inference
import benchmark
import train
import dataset
import model as model_utils
import experiments
import reporting

REAL_PLOT_CURVES = train.plot_curves
REAL_PLOT_MATRICES = experiments.plot_final_matrices

torch.set_num_threads(1)


class TinyModel(nn.Module):
    def __init__(self):
        super().__init__()
        self.conv = nn.Conv2d(3, 4, 1)
        self.bn = nn.BatchNorm2d(4)
        self.head = nn.Linear(4, 9)
        self.pretrained_cfg = {"architecture": "tiny", "tag": "test"}

    def get_classifier(self):
        return self.head

    def forward(self, x):
        return self.head(self.bn(self.conv(x)).mean((2, 3)))


class TechniquesTest(unittest.TestCase):
    def test_losses_and_weights(self):
        logits, labels = torch.randn(12, 9), torch.arange(12) % 9
        ce = nn.CrossEntropyLoss()(logits, labels)
        self.assertTrue(torch.allclose(losses.FocalLoss(0)(logits, labels), ce, atol=1e-6))
        self.assertTrue(torch.allclose(losses.LabelSmoothingCE(0)(logits, labels), ce, atol=1e-6))
        weights = losses.class_weights([1, 2, 4])
        self.assertAlmostEqual(weights.mean().item(), 1.0, places=6)
        self.assertGreater(weights[0].item(), weights[2].item())
        for kind in ("ce", "ls", "focal", "ce_weighted"):
            value = losses.build_criterion(kind, weight=torch.ones(9))(logits, labels)
            self.assertTrue(torch.isfinite(value))

    def test_cutmix_uses_actual_area_and_preserves_input(self):
        images = torch.stack([torch.zeros(3, 10, 10), torch.ones(3, 10, 10)])
        labels = torch.tensor([0, 1])
        with patch("numpy.random.beta", return_value=0.5), patch("numpy.random.randint", return_value=0), \
             patch("torch.randperm", return_value=torch.tensor([1, 0])):
            mixed, (a, b, lam) = losses.mix_batch(images, labels, mode="cutmix")
        self.assertAlmostEqual(1 - lam, mixed[0, 0].mean().item(), places=6)
        self.assertTrue(torch.equal(images[0], torch.zeros_like(images[0])))
        self.assertTrue(torch.equal(a, labels))
        self.assertTrue(torch.equal(b, labels.flip(0)))
        logits = torch.randn(2, 9)
        expected = lam * nn.CrossEntropyLoss()(logits, a) + (1 - lam) * nn.CrossEntropyLoss()(logits, b)
        self.assertTrue(torch.allclose(losses.mixed_loss(nn.CrossEntropyLoss(), logits, (a, b, lam)), expected))

    def test_optimizer_groups_and_ema(self):
        model = TinyModel()
        groups = model_utils.param_groups(model, 1e-4, 1e-3, 0.05)
        found = [p for group in groups for p in group["params"]]
        self.assertEqual(len({id(p) for p in found}), len(list(model.parameters())))
        for group in groups:
            for parameter in group["params"]:
                if parameter.ndim <= 1:
                    self.assertEqual(group["weight_decay"], 0)
        ema = train.EMA(model, 0.5)
        initial = ema.model.head.weight.clone()
        with torch.no_grad():
            model.head.weight.add_(2)
            model.bn.num_batches_tracked.add_(7)
        ema.update(model)
        self.assertTrue(torch.allclose(ema.model.head.weight, initial + 1))
        self.assertEqual(ema.model.bn.num_batches_tracked.item(), 7)
        self.assertFalse(ema.model.training)

    def test_views_temperature_and_bn(self):
        x = torch.arange(3 * 8 * 8).reshape(1, 3, 8, 8).float()
        crops = inference.views_multicrop(x, 6)
        self.assertEqual(len(crops), 5)
        self.assertTrue(torch.equal(crops[0], x[..., :6, :6]))
        self.assertTrue(torch.equal(crops[-1], x[..., 1:7, 1:7]))
        logits = np.tile([8., 0., -1.], (30, 1))
        labels = np.arange(30) % 3
        temperature = inference.fit_temperature(logits, labels)
        before = inference.apply_temperature(logits, 1)
        after = inference.apply_temperature(logits, temperature)
        self.assertTrue(np.allclose(after.sum(1), 1))
        self.assertTrue(np.array_equal(before.argmax(1), after.argmax(1)))
        before_nll = -np.log(before[np.arange(30), labels]).mean()
        after_nll = -np.log(after[np.arange(30), labels]).mean()
        self.assertLessEqual(after_nll, before_nll)
        for space in ("prob", "logit"):
            self.assertTrue(np.allclose(inference.aggregate_views([logits, logits], space).sum(1), 1))
        net = nn.Sequential(nn.Conv2d(3, 4, 3, padding=1), nn.BatchNorm2d(4), nn.ReLU()).eval()
        fused = inference.fuse_conv_bn(net)
        # Kiểm tra trên ảnh có thang giá trị thực tế, tránh sai số do đầu vào quá lớn.
        torch.testing.assert_close(net(x / 255), fused(x / 255), atol=1e-5, rtol=1e-5)
        self.assertIsInstance(net[1], nn.BatchNorm2d)
        self.assertIsInstance(fused[1], nn.Identity)

    def test_benchmark_and_config(self):
        events = []
        stats = benchmark.bench(lambda: events.append("run"), 10, 50, lambda: events.append("sync"))
        self.assertEqual(events.count("run"), 60)
        self.assertEqual(events.count("sync"), 100)
        self.assertLessEqual(stats["p50"], stats["p95"])
        self.assertLessEqual(stats["p95"], stats["p99"])
        overrides = train.parse_overrides(["seed=2", "amp=false", "ema_decay=none"])
        self.assertEqual(overrides, {"seed": 2, "amp": False, "ema_decay": None})
        with self.assertRaises(ValueError):
            benchmark.bench(lambda: None, iters=1)


class WorkflowTest(unittest.TestCase):
    def setUp(self):
        self.old_dir = Path.cwd()
        self.temporary = tempfile.TemporaryDirectory()
        os.chdir(self.temporary.name)
        self.addCleanup(self.temporary.cleanup)
        self.addCleanup(os.chdir, self.old_dir)
        Path("labels").mkdir()
        self.frames = []
        for split in ("train", "val", "test"):
            frame = pd.DataFrame({"Filename": [f"{split}_{i}.jpg" for i in range(9)], "Label": range(9)})
            frame.to_csv(f"labels/{split}_subset0.csv", index=False)
            self.frames.append(frame)
        self.base = train.Config(backbone="tiny.test", epochs=1, batch_size=9, img_size=16,
                                 labels_dir="labels", num_workers=0, amp=False)
        experiments.write_json("training_selection.json", {
            "baseline_config": asdict(replace(self.base, exp_id="T00")),
            "best_config": asdict(replace(self.base, exp_id="T01")),
        })
        experiments.write_json("inference_selection.json", {
            "training_config": asdict(replace(self.base, exp_id="T01")), "method": "hflip_prob",
        })
        self.predict_calls = []
        original_predict = inference.predict_probs
        def record_predict(model, loader, *args, **kw):
            self.predict_calls.append(loader[0][2][0].split("_")[0])
            return original_predict(model, loader, *args, **kw)
        self.patches = [
            patch.object(dataset, "load_split", return_value=tuple(self.frames)),
            patch.object(dataset, "check_split", return_value={}),
            patch.object(dataset, "make_loader", side_effect=lambda df, *a, **kw: self.loader(df)),
            patch.object(model_utils, "build_model", side_effect=lambda *a, **kw: TinyModel()),
            patch.object(model_utils, "count_gmacs", return_value=0.001),
            patch.object(train, "plot_curves"), patch.object(experiments, "plot_final_matrices"),
            patch.object(torch.cuda, "is_available", return_value=False),
            patch.object(experiments, "make_eval_loader", side_effect=lambda cfg, df, method: self.loader(df)),
            patch.object(experiments, "measure_method", return_value=[{
                "p50": 1., "p95": 2., "p99": 3., "batch": 1, "images_per_s": 1000., "gpu": "CPU"}]),
            patch.object(inference, "predict_probs", side_effect=record_predict),
        ]
        for item in self.patches:
            item.start()

    def loader(self, frame):
        return [(torch.ones(len(frame), 3, 16, 16), torch.tensor(frame.Label.to_numpy()), frame.Filename.tolist())]

    def tearDown(self):
        for item in reversed(self.patches):
            item.stop()
        os.chdir(self.old_dir)
        self.temporary.cleanup()

    def test_full_final_loop_cache_and_lock(self):
        table = experiments.final_step()
        self.assertEqual(self.predict_calls.count("test"), 6)
        self.assertEqual(len(table.loc[table.kind.eq("seed")]), 6)
        self.assertEqual(len(table.loc[table.kind.eq("mean_std")]), 2)
        self.assertTrue(Path("predictions/F01_seed2_test.csv").is_file())
        self.assertTrue(Path("predictions/F01uncal_seed2_test.csv").is_file())
        self.assertTrue(Path("runs/T00/seed2/best.pt").is_file())
        self.assertTrue(Path("results.xlsx").is_file())
        calls_before = len(self.predict_calls)
        experiments.final_step()
        self.assertEqual(len(self.predict_calls), calls_before)
        selection = experiments.read_json("inference_selection.json")
        selection["method"] = "identity"
        experiments.write_json("inference_selection.json", selection)
        with self.assertRaises(ValueError):
            experiments.final_step()

    def test_training_axes_combination_and_inference(self):
        # Toàn bộ 10 cấu hình train chạy bằng TinyModel, gồm CutMix và EMA.
        result = experiments.run_cached(replace(self.base, exp_id="B01"))
        pd.DataFrame([result]).to_csv("backbones.csv", index=False)
        experiments.write_json("backbone_selection.json", {"selected_backbones": [self.base.backbone]})
        table = experiments.training_step("images", "labels")
        self.assertEqual(len(table), 10)
        self.assertTrue(set(table.axis) >= {"B", "C", "D", "F", "combination"})
        self.assertEqual(self.predict_calls.count("test"), 0)
        base_config = experiments.read_json("runs/T00/seed0/config.json")
        for exp_id in range(1, 9):
            cfg = experiments.read_json(f"runs/T{exp_id:02d}/seed0/config.json")
            # Nhóm kỹ thuật có thể cần thêm tham số đi kèm, nhưng không đổi recipe nền.
            for key in ("backbone", "seed", "img_size", "epochs", "batch_size", "lr_backbone", "lr_head"):
                self.assertEqual(cfg[key], base_config[key])
        inference_table = experiments.inference_step(batches=(1,))
        self.assertEqual(len(inference_table), 5)
        self.assertEqual(self.predict_calls.count("test"), 0)
        self.assertTrue(Path("curves/inference_tradeoff.png").exists())
        self.assertTrue(Path("inference_selection.json").exists())
        calibrated = inference_table.loc[inference_table.exp_id.eq("I04")].iloc[0]
        baseline = inference_table.loc[inference_table.exp_id.eq("I00")].iloc[0]
        self.assertAlmostEqual(calibrated.val_macro_f1, baseline.val_macro_f1)

    def test_interrupted_test_is_not_repeated(self):
        cfg = replace(self.base, exp_id="F01")
        folder = train.run_dir(cfg)
        folder.mkdir(parents=True)
        experiments.write_json(folder / "test_started.json", {})
        with self.assertRaises(RuntimeError):
            experiments.final_prediction(cfg, "identity", torch.device("cpu"))
        self.assertEqual(self.predict_calls, [])

    def test_deliverables_verify_export_and_preserve_manual_report(self):
        # Tạo đủ bằng chứng nhỏ của 5 backbone, 10 recipe và 3 seed chung kết.
        backbone_rows = [experiments.run_cached(replace(self.base, exp_id=f"B{i:02d}")) for i in range(1, 6)]
        pd.DataFrame(backbone_rows).to_csv("backbones.csv", index=False)
        experiments.save_sheet(pd.DataFrame(backbone_rows), "Backbones")
        experiments.write_json("backbone_selection.json", {"selected_backbones": [self.base.backbone]})
        experiments.training_step("images", "labels")
        def reports(model, cfg, device, method, temperature=1.0, batches=(1, 32), iters=50):
            return [{"p50": 1., "p95": 2., "p99": 3., "mean": 1.5, "n": iters,
                     "batch": batch, "images_per_s": batch * 1000., "gpu": "CPU", "dtype": "fp32",
                     "img_size": cfg.img_size, "input_size": cfg.img_size, "fused_bn": False,
                     "torch": torch.__version__, "method": method, "preprocessing": "excluded"} for batch in batches]
        with patch.object(experiments, "measure_method", side_effect=reports):
            experiments.inference_step()
            experiments.final_step()
        experiments.save_sheet(pd.DataFrame({"note": ["keep this sheet"]}), "UserNotes")
        before = list(self.predict_calls)
        with patch.object(train, "plot_curves", new=REAL_PLOT_CURVES), \
             patch.object(experiments, "plot_final_matrices", new=REAL_PLOT_MATRICES):
            products = reporting.deliverables_step(submission_dir="delivery")
            self.assertEqual(before, self.predict_calls)  # Không forward model nữa.
            delivered = Path(products["submission_dir"])
            self.assertEqual(products["verified_runs"], 20)
            self.assertTrue((delivered / "report.md").is_file())
            self.assertTrue((delivered / "README.md").is_file())
            self.assertTrue((delivered / "eval_out/F01_summary.json").is_file())
            manifest = experiments.read_json(delivered / "artifact_manifest.json")
            self.assertIn("results.xlsx", manifest)
            self.assertFalse(any(name.endswith(".pt") for name in manifest))
            import zipfile
            with zipfile.ZipFile(products["archive"]) as bundle:
                names = bundle.namelist()
                self.assertIn("results.xlsx", names)
                self.assertIn("artifact_manifest.json", names)
                self.assertFalse(any(name.endswith(".pt") or "/data/" in name for name in names))
            from openpyxl import load_workbook
            with pd.ExcelFile(delivered / "results.xlsx") as book:
                self.assertTrue(set(reporting.SHEETS) <= set(book.sheet_names))
                self.assertIn("UserNotes", book.sheet_names)
            book = load_workbook(delivered / "results.xlsx")
            self.assertEqual(book["Summary"].freeze_panes, "A2")
            self.assertEqual(book["Summary"].max_row, 11)
            book.close()
            (delivered / "report.md").write_text("My reviewed report", encoding="utf-8")
            # Lần xuất tiếp theo vẫn giữ nguyên báo cáo người dùng đã sửa.
            with patch.object(train, "plot_curves"), patch.object(experiments, "plot_final_matrices"):
                reporting.deliverables_step(submission_dir="delivery")
            self.assertEqual((delivered / "report.md").read_text(encoding="utf-8"), "My reviewed report")
        # Khi bảng bị chỉnh lệch CSV, từ chối xuất trước khi ghi báo cáo.
        tables = pd.read_excel("results.xlsx", sheet_name=None)
        tables["Backbones"].loc[0, "val_macro_f1"] = 0.99
        experiments.save_sheet(tables["Backbones"], "Backbones")
        with self.assertRaisesRegex(ValueError, "khớp"):
            reporting.deliverables_step(submission_dir="delivery")
        self.assertEqual(before, self.predict_calls)


if __name__ == "__main__":
    unittest.main()
